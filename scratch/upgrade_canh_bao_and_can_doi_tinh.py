import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from datetime import datetime, timedelta

paths = [
    r'c:\Users\lap4all\Desktop\BAI_NOP_CAPACITY_GHN\BAO_CAO_CAPACITY_11.11_FINAL.xlsx',
    r'c:\Users\lap4all\Desktop\New folder\Capacity_Planning_Test\[GHN]_Capacity Planning_Test_T10.2026\BAO_CAO_CAPACITY_11.11_FINAL.xlsx'
]

# Load lay_hang data
wb_lay = openpyxl.load_workbook(r'c:\Users\lap4all\Desktop\New folder\Capacity_Planning_Test\[GHN]_Capacity Planning_Test_T10.2026\lay_hang.xlsx', data_only=True)
ws_lay_ngay = wb_lay['lay_ngay']
lay_dict = {}
for r in range(2, ws_lay_ngay.max_row + 1):
    d = str(ws_lay_ngay.cell(r, 1).value)[:10]
    bc = ws_lay_ngay.cell(r, 3).value
    lay = ws_lay_ngay.cell(r, 6).value
    lay_dict[(d, bc)] = lay

# Styling definitions
navy_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
slate_fill = PatternFill(start_color="204A6E", end_color="204A6E", fill_type="solid")
white_bold_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
thin_border = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)
header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
regular_font = Font(name="Segoe UI", size=10)
bold_font = Font(name="Segoe UI", size=10, bold=True)

# Post offices grouping
tay_lam_bcs = [
    ('Bến Sỏi', 'C'),
    ('Cầu Mây', 'D'),
    ('Gò Đào', 'E'),
    ('Lâm Viên 1', 'F'),
    ('Lâm Viên 2', 'G'),
    ('Rạch Sen', 'H')
]

bac_hoa_bcs = [
    ('CK Bắc Hòa', 'I'),
    ('Đồng Tre', 'J'),
    ('Hòa Lộc', 'K'),
    ('Kênh Đào', 'L'),
    ('Phú Điền', 'M'),
    ('Suối Ngọc', 'N')
]

dates = [datetime(2026, 11, 8) + timedelta(days=i) for i in range(14)]

for p in paths:
    print('Processing:', p)
    wb = openpyxl.load_workbook(p)
    
    # -------------------------------------------------------------------------
    # 1. UPGRADE SHEET canh_bao (Cách 1: 2 Khối Tỉnh riêng biệt)
    # -------------------------------------------------------------------------
    ws_cb = wb['canh_bao']
    
    # Clear existing rows from 8 onwards
    for r in range(8, ws_cb.max_row + 1):
        for c in range(1, 13):
            ws_cb.cell(r, c).value = None
            ws_cb.cell(r, c).fill = PatternFill(fill_type=None)
            ws_cb.cell(r, c).border = Border()
            
    for m in list(ws_cb.merged_cells.ranges):
        if m.min_row >= 8:
            ws_cb.unmerge_cells(str(m))
            
    # Row 8: Table Header
    headers_cb = ['Ngày', 'Bưu Cục', 'Tỉnh', 'Đơn Mới Về', 'Cap Giao Thực Tế', 'Chênh Lệch Giao', '% Vượt Giao', 'Cảnh Báo Chiều Giao', 'Đơn Lấy Dự Báo', 'Trần Lấy TB', 'Chênh Lệch Lấy']
    for c, h in enumerate(headers_cb, 1):
        cell = ws_cb.cell(8, c, h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = thin_border
    ws_cb.row_dimensions[8].height = 28
    
    # Row 9: Section Header TÂY LÂM
    ws_cb.merge_cells('A9:K9')
    cell_tl_sec = ws_cb.cell(9, 1, 'TỈNH TÂY LÂM (6 BƯU CỤC - THEO DÕI 14 NGÀY TỪ 08/11 ĐẾN 21/11/2026)')
    cell_tl_sec.font = white_bold_font
    cell_tl_sec.fill = navy_fill
    cell_tl_sec.alignment = Alignment(horizontal='center', vertical='center')
    ws_cb.row_dimensions[9].height = 26
    for c in range(1, 12):
        ws_cb.cell(9, c).fill = navy_fill
        ws_cb.cell(9, c).border = thin_border
        
    cur_row = 10
    # Rows 10 to 93: Tây Lâm data (14 days x 6 bcs = 84 rows)
    for d_idx, dt in enumerate(dates):
        d_str = dt.strftime('%Y-%m-%d')
        d_disp = dt.strftime('%d/%m/%Y')
        du_bao_row = 9 + d_idx
        
        for bc_name, col_letter in tay_lam_bcs:
            r = cur_row
            ws_cb.row_dimensions[r].height = 20
            
            # A: Ngay
            c_a = ws_cb.cell(r, 1, d_disp)
            c_a.alignment = Alignment(horizontal='center', vertical='center')
            
            # B: Buu Cuc
            c_b = ws_cb.cell(r, 2, bc_name)
            c_b.alignment = Alignment(horizontal='left', vertical='center')
            c_b.font = bold_font
            
            # C: Tinh
            c_c = ws_cb.cell(r, 3, 'Tây Lâm')
            c_c.alignment = Alignment(horizontal='center', vertical='center')
            
            # D: Don Moi Ve
            c_d = ws_cb.cell(r, 4, f'=du_bao_giao!{col_letter}{du_bao_row}')
            c_d.alignment = Alignment(horizontal='right', vertical='center')
            c_d.number_format = '#,##0'
            
            # E: Cap Giao Thuc Te
            c_e = ws_cb.cell(r, 5, f'=VLOOKUP(B{r}, cap_hieu_chinh!$A$5:$H$18, 8, FALSE)')
            c_e.alignment = Alignment(horizontal='right', vertical='center')
            c_e.number_format = '#,##0'
            
            # F: Chenh Lech Giao
            c_f = ws_cb.cell(r, 6, f'=D{r}-E{r}')
            c_f.alignment = Alignment(horizontal='right', vertical='center')
            c_f.number_format = '#,##0'
            
            # G: % Vuot Giao
            c_g = ws_cb.cell(r, 7, f'=F{r}/E{r}')
            c_g.alignment = Alignment(horizontal='right', vertical='center')
            c_g.number_format = '0.0%'
            
            # H: Canh Bao Chieu Giao
            c_h = ws_cb.cell(r, 8, f'=IF(AND(F{r}>=50, G{r}>=0.15), "🔴 Chắc vượt", IF(AND(F{r}>=50, G{r}>0), "🟡 Có thể vượt", "🟢 An toàn"))')
            c_h.alignment = Alignment(horizontal='center', vertical='center')
            
            # I: Don Lay Du Bao
            lay_val = lay_dict.get((d_str, bc_name), '')
            c_i = ws_cb.cell(r, 9, lay_val)
            c_i.alignment = Alignment(horizontal='right', vertical='center')
            if isinstance(lay_val, (int, float)):
                c_i.number_format = '#,##0'
                
            # J: Tran Lay TB
            c_j = ws_cb.cell(r, 10, f'=VLOOKUP(B{r}, chat_luong_lay!$A$5:$C$18, 3, FALSE)')
            c_j.alignment = Alignment(horizontal='right', vertical='center')
            c_j.number_format = '#,##0'
            
            # K: Chenh Lech Lay
            if bc_name == 'Lâm Viên 2':
                c_k = ws_cb.cell(r, 11, '-')
                c_k.alignment = Alignment(horizontal='center', vertical='center')
            else:
                c_k = ws_cb.cell(r, 11, f'=IF(ISNUMBER(I{r}), I{r}-J{r}, "-")')
                c_k.alignment = Alignment(horizontal='right', vertical='center')
                c_k.number_format = '#,##0'
                
            for c in range(1, 12):
                ws_cb.cell(r, c).font = regular_font
                ws_cb.cell(r, c).border = thin_border
            c_b.font = bold_font
            
            cur_row += 1

    # Row 94: Section Header BẮC HÒA
    row_bh_sec = cur_row
    ws_cb.merge_cells(f'A{row_bh_sec}:K{row_bh_sec}')
    cell_bh_sec = ws_cb.cell(row_bh_sec, 1, 'TỈNH BẮC HÒA (6 BƯU CỤC - THEO DÕI 14 NGÀY TỪ 08/11 ĐẾN 21/11/2026)')
    cell_bh_sec.font = white_bold_font
    cell_bh_sec.fill = slate_fill
    cell_bh_sec.alignment = Alignment(horizontal='center', vertical='center')
    ws_cb.row_dimensions[row_bh_sec].height = 26
    for c in range(1, 12):
        ws_cb.cell(row_bh_sec, c).fill = slate_fill
        ws_cb.cell(row_bh_sec, c).border = thin_border
        
    cur_row += 1
    # Rows 95 to 178: Bắc Hòa data (14 days x 6 bcs = 84 rows)
    for d_idx, dt in enumerate(dates):
        d_str = dt.strftime('%Y-%m-%d')
        d_disp = dt.strftime('%d/%m/%Y')
        du_bao_row = 9 + d_idx
        
        for bc_name, col_letter in bac_hoa_bcs:
            r = cur_row
            ws_cb.row_dimensions[r].height = 20
            
            # A: Ngay
            c_a = ws_cb.cell(r, 1, d_disp)
            c_a.alignment = Alignment(horizontal='center', vertical='center')
            
            # B: Buu Cuc
            c_b = ws_cb.cell(r, 2, bc_name)
            c_b.alignment = Alignment(horizontal='left', vertical='center')
            c_b.font = bold_font
            
            # C: Tinh
            c_c = ws_cb.cell(r, 3, 'Bắc Hòa')
            c_c.alignment = Alignment(horizontal='center', vertical='center')
            
            # D: Don Moi Ve
            c_d = ws_cb.cell(r, 4, f'=du_bao_giao!{col_letter}{du_bao_row}')
            c_d.alignment = Alignment(horizontal='right', vertical='center')
            c_d.number_format = '#,##0'
            
            # E: Cap Giao Thuc Te
            c_e = ws_cb.cell(r, 5, f'=VLOOKUP(B{r}, cap_hieu_chinh!$A$5:$H$18, 8, FALSE)')
            c_e.alignment = Alignment(horizontal='right', vertical='center')
            c_e.number_format = '#,##0'
            
            # F: Chenh Lech Giao
            c_f = ws_cb.cell(r, 6, f'=D{r}-E{r}')
            c_f.alignment = Alignment(horizontal='right', vertical='center')
            c_f.number_format = '#,##0'
            
            # G: % Vuot Giao
            c_g = ws_cb.cell(r, 7, f'=F{r}/E{r}')
            c_g.alignment = Alignment(horizontal='right', vertical='center')
            c_g.number_format = '0.0%'
            
            # H: Canh Bao Chieu Giao
            c_h = ws_cb.cell(r, 8, f'=IF(AND(F{r}>=50, G{r}>=0.15), "🔴 Chắc vượt", IF(AND(F{r}>=50, G{r}>0), "🟡 Có thể vượt", "🟢 An toàn"))')
            c_h.alignment = Alignment(horizontal='center', vertical='center')
            
            # I: Don Lay Du Bao
            lay_val = lay_dict.get((d_str, bc_name), '')
            c_i = ws_cb.cell(r, 9, lay_val)
            c_i.alignment = Alignment(horizontal='right', vertical='center')
            if isinstance(lay_val, (int, float)):
                c_i.number_format = '#,##0'
                
            # J: Tran Lay TB
            c_j = ws_cb.cell(r, 10, f'=VLOOKUP(B{r}, chat_luong_lay!$A$5:$C$18, 3, FALSE)')
            c_j.alignment = Alignment(horizontal='right', vertical='center')
            c_j.number_format = '#,##0'
            
            # K: Chenh Lech Lay
            c_k = ws_cb.cell(r, 11, f'=IF(ISNUMBER(I{r}), I{r}-J{r}, "-")')
            c_k.alignment = Alignment(horizontal='right', vertical='center')
            c_k.number_format = '#,##0'
                
            for c in range(1, 12):
                ws_cb.cell(r, c).font = regular_font
                ws_cb.cell(r, c).border = thin_border
            c_b.font = bold_font
            
            cur_row += 1

    # Widths for canh_bao
    ws_cb.column_dimensions['A'].width = 13
    ws_cb.column_dimensions['B'].width = 16
    ws_cb.column_dimensions['C'].width = 13
    ws_cb.column_dimensions['D'].width = 14
    ws_cb.column_dimensions['E'].width = 17
    ws_cb.column_dimensions['F'].width = 17
    ws_cb.column_dimensions['G'].width = 13
    ws_cb.column_dimensions['H'].width = 18
    ws_cb.column_dimensions['I'].width = 16
    ws_cb.column_dimensions['J'].width = 14
    ws_cb.column_dimensions['K'].width = 16

    # -------------------------------------------------------------------------
    # 2. UPGRADE SHEET can_doi_tinh WITH PROVINCE SECTION MERGERS
    # -------------------------------------------------------------------------
    ws_cd = wb['can_doi_tinh']
    
    # Detail table items
    tay_lam_plan = [
        ('Bến Sỏi', 'Tây Lâm', '=du_bao_giao!C13', '=VLOOKUP(A{r}, cap_hieu_chinh!$A$5:$H$18, 8, FALSE)', '=C{r}-D{r}', 'BỔ SUNG 6 THỜI VỤ (+252 đơn/ngày) + Tuyển bù nhân sự', '10/11', 'Bàn giao 2 phường giáp ranh cho Gò Đào nhận tải giao thay'),
        ('Cầu Mây', 'Tây Lâm', '=du_bao_giao!D13', '=VLOOKUP(A{r}, cap_hieu_chinh!$A$5:$H$18, 8, FALSE)', '=C{r}-D{r}', 'BỔ SUNG 4 THỜI VỤ (+168 đơn/ngày) hỗ trợ lõi thị trấn', '10/11', 'Gán Ca 1 sớm trước 8h + khoán thêm đơn cho rider cứng'),
        ('Lâm Viên 1', 'Tây Lâm', '=du_bao_giao!F13', '=VLOOKUP(A{r}, cap_hieu_chinh!$A$5:$H$18, 8, FALSE)', '=C{r}-D{r}', 'BỔ SUNG 2 THỜI VỤ (+84 đơn/ngày) phụ tuyến nội thành', '10/11', 'Lâm Viên 2 hỗ trợ giao chéo các khu vực giáp ranh'),
        ('Lâm Viên 2', 'Tây Lâm', '=du_bao_giao!G13', '=VLOOKUP(A{r}, cap_hieu_chinh!$A$5:$H$18, 8, FALSE)', '=C{r}-D{r}', 'KHÔNG THÊM NGƯỜI (Rider cứng kèm cặp rider mới)', '08/11', 'Lâm Viên 1 gánh lại một số tuyến vùng ven ngày 11-12/11'),
        ('Rạch Sen', 'Tây Lâm', '=du_bao_giao!H13', '=VLOOKUP(A{r}, cap_hieu_chinh!$A$5:$H$18, 8, FALSE)', '=C{r}-D{r}', 'TỰ CÂN ĐỐI (Vượt nhẹ 6.9%, tăng năng suất +5 đơn/người)', '11/11', 'Không cần thêm người, tự điều phối nội bộ'),
        ('Gò Đào', 'Tây Lâm', '=du_bao_giao!E13', '=VLOOKUP(A{r}, cap_hieu_chinh!$A$5:$H$18, 8, FALSE)', '=C{r}-D{r}', 'ĐIỂM ĐỆM NHẬN TẢI (Dư 537 đơn): Điều 2 shipper sang Bến Sỏi', '10/11', 'Nhận 2 phường giáp ranh của Bến Sỏi gánh phụ toàn diện')
    ]
    
    bac_hoa_plan = [
        ('CK Bắc Hòa', 'Bắc Hòa', '=du_bao_giao!I13', '=VLOOKUP(A{r}, cap_hieu_chinh!$A$5:$H$18, 8, FALSE)', '=C{r}-D{r}', 'Thêm xe tải nhỏ/xe ba gác (Hàng cồng kềnh xe máy không chở được)', '10/11', 'Hẹn giờ giao theo khung với khách lớn, tránh dồn hàng'),
        ('Đồng Tre', 'Bắc Hòa', '=du_bao_giao!J13', '=VLOOKUP(A{r}, cap_hieu_chinh!$A$5:$H$18, 8, FALSE)', '=C{r}-D{r}', 'TỰ CÂN ĐỐI NỘI BỘ (Vượt nhẹ 37 đơn ngày 12/11)', '11/11', 'Tăng tỷ lệ gán Ca 1 >85%, tự giải phóng trong ngày'),
        ('Hòa Lộc', 'Bắc Hòa', '=du_bao_giao!K13', '=VLOOKUP(A{r}, cap_hieu_chinh!$A$5:$H$18, 8, FALSE)', '=C{r}-D{r}', 'Giao đủ cap, NGHẼN CHIỀU LẤY (OPR tụt 65%, đơn lấy vượt 81%)', '09/11', 'KTC ĐIỀU XE TẢI BỐC THẲNG TẠI 3 SHOP LỚN (Direct Pickup)'),
        ('Kênh Đào', 'Bắc Hòa', '=du_bao_giao!L13', '=VLOOKUP(A{r}, cap_hieu_chinh!$A$5:$H$18, 8, FALSE)', '=C{r}-D{r}', 'DƯ CÔNG SUẤT: Ghép tuyến hỗ trợ Suối Ngọc và Phú Điền', '09/11', 'Hỗ trợ lấy hàng cho Hòa Lộc các shop vùng ven giáp ranh'),
        ('Phú Điền', 'Bắc Hòa', '=du_bao_giao!M13', '=VLOOKUP(A{r}, cap_hieu_chinh!$A$5:$H$18, 8, FALSE)', '=C{r}-D{r}', 'THÊM NGƯỜI KHÔNG GIẢI QUYẾT ĐƯỢC (Kẹt sàn 90m2 & xe 10h)', 'Trước 11/11', 'YÊU CẦU KTC ĐỔI XE VỀ 6H30 SÁNG; Thuê bãi đệm dã chiến ngoài kho'),
        ('Suối Ngọc', 'Bắc Hòa', '=du_bao_giao!N13', '=VLOOKUP(A{r}, cap_hieu_chinh!$A$5:$H$18, 8, FALSE)', '=C{r}-D{r}', 'DƯ CÔNG SUẤT: Phối hợp Kênh Đào giữ vững địa bàn', '09/11', 'Chia sẻ đầu xe lấy hàng nếu KTC yêu cầu hỗ trợ')
    ]
    
    # Update Part 1 summary formulas
    # Row 6: Tây Lâm will be rows 12 to 17
    ws_cd['B6'] = '=SUM(C12:C17)'
    ws_cd['D6'] = '=SUM(D12:D17)'
    ws_cd['E6'] = '=B6-D6'
    ws_cd['F6'] = '=E6/D6'
    
    # Row 7: Bắc Hòa will be rows 19 to 24
    ws_cd['B7'] = '=SUM(C19:C24)'
    ws_cd['D7'] = '=SUM(D19:D24)'
    ws_cd['E7'] = '=B7-D7'
    ws_cd['F7'] = '=E7/D7'
    
    # Clear Part 2 rows 10 to 30
    for r in range(10, 30):
        for c in range(1, 10):
            ws_cd.cell(r, c).value = None
            ws_cd.cell(r, c).fill = PatternFill(fill_type=None)
            ws_cd.cell(r, c).border = Border()
    for m in list(ws_cd.merged_cells.ranges):
        if m.min_row >= 10:
            ws_cd.unmerge_cells(str(m))
            
    # Row 10: Table Header
    headers_cd = ['Bưu Cục', 'Tỉnh', 'Đơn Peak 12/11', 'Cap Thực Tế', 'Chênh Lệch', 'Phương Án Phân Bổ Nhân Sự', 'Ngày Bắt Đầu', 'Giải Pháp Thay Thế Nếu Không Thêm Người']
    for c, h in enumerate(headers_cd, 1):
        cell = ws_cd.cell(10, c, h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = thin_border
    ws_cd.row_dimensions[10].height = 28
    
    # Row 11: Section Header TÂY LÂM
    ws_cd.merge_cells('A11:H11')
    cell_cd_tl = ws_cd.cell(11, 1, 'TỈNH TÂY LÂM (6 BƯU CỤC - THIẾU SỨC GIAO THỰC SỰ +1.091 ĐƠN PEAK)')
    cell_cd_tl.font = white_bold_font
    cell_cd_tl.fill = navy_fill
    cell_cd_tl.alignment = Alignment(horizontal='center', vertical='center')
    ws_cd.row_dimensions[11].height = 24
    for c in range(1, 9):
        ws_cd.cell(11, c).fill = navy_fill
        ws_cd.cell(11, c).border = thin_border
        
    # Rows 12 to 17: Tây Lâm detail
    for idx, item in enumerate(tay_lam_plan):
        r = 12 + idx
        ws_cd.row_dimensions[r].height = 22
        vals = [
            item[0], item[1], item[2], item[3].format(r=r), item[4].format(r=r),
            item[5], item[6], item[7]
        ]
        for c, v in enumerate(vals, 1):
            cell = ws_cd.cell(r, c, v)
            cell.font = regular_font
            cell.border = thin_border
            if c in [1, 2, 7]:
                cell.alignment = Alignment(horizontal='center', vertical='center')
            elif c in [3, 4, 5]:
                cell.alignment = Alignment(horizontal='right', vertical='center')
                cell.number_format = '#,##0'
            else:
                cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
            if c == 1:
                cell.font = bold_font

    # Row 18: Section Header BẮC HÒA
    ws_cd.merge_cells('A18:H18')
    cell_cd_bh = ws_cd.cell(18, 1, 'TỈNH BẮC HÒA (6 BƯU CỤC - TỔNG THỂ ĐỦ SỨC GIAO, NGHẼN HẠ TẦNG TẠI PHÚ ĐIỀN)')
    cell_cd_bh.font = white_bold_font
    cell_cd_bh.fill = slate_fill
    cell_cd_bh.alignment = Alignment(horizontal='center', vertical='center')
    ws_cd.row_dimensions[18].height = 24
    for c in range(1, 9):
        ws_cd.cell(18, c).fill = slate_fill
        ws_cd.cell(18, c).border = thin_border

    # Rows 19 to 24: Bắc Hòa detail
    for idx, item in enumerate(bac_hoa_plan):
        r = 19 + idx
        ws_cd.row_dimensions[r].height = 22
        vals = [
            item[0], item[1], item[2], item[3].format(r=r), item[4].format(r=r),
            item[5], item[6], item[7]
        ]
        for c, v in enumerate(vals, 1):
            cell = ws_cd.cell(r, c, v)
            cell.font = regular_font
            cell.border = thin_border
            if c in [1, 2, 7]:
                cell.alignment = Alignment(horizontal='center', vertical='center')
            elif c in [3, 4, 5]:
                cell.alignment = Alignment(horizontal='right', vertical='center')
                cell.number_format = '#,##0'
            else:
                cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
            if c == 1:
                cell.font = bold_font

    # Column dimensions for can_doi_tinh
    ws_cd.column_dimensions['A'].width = 16
    ws_cd.column_dimensions['B'].width = 14
    ws_cd.column_dimensions['C'].width = 16
    ws_cd.column_dimensions['D'].width = 16
    ws_cd.column_dimensions['E'].width = 16
    ws_cd.column_dimensions['F'].width = 38
    ws_cd.column_dimensions['G'].width = 14
    ws_cd.column_dimensions['H'].width = 45

    wb.save(p)
    print('Completed upgrade for:', p)

