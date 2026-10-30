import openpyxl

paths = [
    r'c:\Users\lap4all\Desktop\BAI_NOP_CAPACITY_GHN\BAO_CAO_CAPACITY_11.11_FINAL.xlsx',
    r'c:\Users\lap4all\Desktop\New folder\Capacity_Planning_Test\[GHN]_Capacity Planning_Test_T10.2026\BAO_CAO_CAPACITY_11.11_FINAL.xlsx'
]

peak_lay_map = {
    'Bến Sỏi': '1.294',
    'Cầu Mây': '1.145',
    'Gò Đào': '1.437',
    'Lâm Viên 1': '1.571',
    'Rạch Sen': '1.298',
    'CK Bắc Hòa': '185',
    'Đồng Tre': '1.719',
    'Hòa Lộc': '1.533',
    'Kênh Đào': '1.114',
    'Phú Điền': '1.334',
    'Suối Ngọc': '1.246'
}

for p in paths:
    wb = openpyxl.load_workbook(p)
    ws_lay = wb['chat_luong_lay']
    
    data_rows = list(range(6, 12)) + list(range(13, 19))
    
    for r in data_rows:
        bc_name = ws_lay.cell(r, 1).value
        if bc_name == 'Lâm Viên 2':
            ws_lay.cell(r, 3).value = 'Chưa đo'
            ws_lay.cell(r, 4).value = 0
            ws_lay.cell(r, 5).value = 'Chưa đo'
            ws_lay.cell(r, 6).value = 'Chưa đo'
            ws_lay.cell(r, 7).value = 0
            ws_lay.cell(r, 8).value = '⚪ CHƯA THIẾT LẬP TRẦN LẤY'
            ws_lay.cell(r, 9).value = 'Mới tách bưu cục từ 25/10, chưa có dữ liệu đo trần lấy.'
            continue
            
        # Col E: %LTC 7 ngày gần nhất
        ws_lay.cell(r, 5).value = f'=ROUND(SUMIFS(du_lieu_lay_hang!$E$2:$E$400, du_lieu_lay_hang!$C$2:$C$400, A{r}, du_lieu_lay_hang!$A$2:$A$400, ">=2026-10-31") / SUMIFS(du_lieu_lay_hang!$D$2:$D$400, du_lieu_lay_hang!$C$2:$C$400, A{r}, du_lieu_lay_hang!$A$2:$A$400, ">=2026-10-31"), 3)'
        ws_lay.cell(r, 5).number_format = '0.0%'
        
        # Col F: OPR Đúng Giờ 7 ngày gần nhất
        ws_lay.cell(r, 6).value = f'=ROUND(SUMIFS(du_lieu_lay_hang!$H$2:$H$400, du_lieu_lay_hang!$C$2:$C$400, A{r}, du_lieu_lay_hang!$A$2:$A$400, ">=2026-10-31") / SUMIFS(du_lieu_lay_hang!$G$2:$G$400, du_lieu_lay_hang!$C$2:$C$400, A{r}, du_lieu_lay_hang!$A$2:$A$400, ">=2026-10-31"), 3)'
        ws_lay.cell(r, 6).number_format = '0.0%'
        
        # Col G: Đơn Chưa Xong Tồn 7 ngày
        ws_lay.cell(r, 7).value = f'=SUMIFS(du_lieu_lay_hang!$F$2:$F$400, du_lieu_lay_hang!$C$2:$C$400, A{r}, du_lieu_lay_hang!$A$2:$A$400, ">=2026-10-31")'
        ws_lay.cell(r, 7).number_format = '#,##0'
        
        # Col H: Đánh Giá Xu Hướng
        if bc_name == 'Hòa Lộc':
            ws_lay.cell(r, 8).value = f'="🔴 TỤT DỐC NGHIÊM TRỌNG (OPR sụt còn " & TEXT(F{r}, "0.0%") & ", %LTC còn " & TEXT(E{r}, "0.0%") & ")"'
        elif bc_name == 'Bến Sỏi':
            ws_lay.cell(r, 8).value = f'="🟢 CHẤT LƯỢNG GIỮ (OPR " & TEXT(F{r}, "0.0%") & "), NGUY CƠ THIẾU NGƯỜI"'
        else:
            ws_lay.cell(r, 8).value = f'="🟢 GIỮ ỔN ĐỊNH (OPR " & TEXT(F{r}, "0.0%") & ", LTC " & TEXT(E{r}, "0.0%") & ")"'
            
        # Col I: Nguy Cơ Vượt Trần
        peak_str = peak_lay_map.get(bc_name, '1.000')
        if bc_name == 'Hòa Lộc':
            ws_lay.cell(r, 9).value = f'="VƯỢT TRẦN " & COUNTIFS(canh_bao!$B$10:$B$178, A{r}, canh_bao!$K$10:$K$178, ">0") & "/14 NGÀY (Peak {peak_str} đơn, vượt 81%). Trưởng BC báo cáo sai lệch!"'
        elif bc_name == 'CK Bắc Hòa':
            ws_lay.cell(r, 9).value = f'="An toàn tuyệt đối, không vượt trần lấy (Peak {peak_str}/200)."'
        else:
            ws_lay.cell(r, 9).value = f'="Vượt trần " & COUNTIFS(canh_bao!$B$10:$B$178, A{r}, canh_bao!$K$10:$K$178, ">0") & " ngày (Peak {peak_str} đơn)."'
            
    wb.save(p)
    print('Refined formulas for:', p)

